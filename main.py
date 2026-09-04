import argparse
from importlib.resources import path
import json
import re
import shutil
import calendar
from pathlib import Path
from datetime import date, datetime

from openpyxl import Workbook, load_workbook

try:
    import winsound
except ImportError:
    winsound = None


CANONICAL_HEADERS = [
    "DATE", "AREA", "CATEGORY", "ENTRY", "AMOUNT (₹)", "PAYMENT / NOTE", "SOURCE TYPE"
]
MILK_HEADERS = [
    "DATE", "SUPPLIER", "MILK QUANTITY", "UNIT", "TRIAL / NOTE", "PAYMENT (₹)"
]


def human_confirmation_alert(message="HUMAN CONFIRMATION REQUIRED"):
    """Alert the user when Purnima Automation needs a human decision."""
    print("\nHUMAN CONFIRMATION REQUIRED")
    print("Press Enter to acknowledge this message.")
    if winsound is not None:
        winsound.Beep(1000, 300)
        winsound.Beep(1000, 300)
        winsound.Beep(1000, 300)

def wait_for_human_confirmation():
    """Pause until a human acknowledges the blocked-run alert."""
    human_confirmation_alert()
    input()


class SafetyGateError(Exception):
    """Raised when the Purnima Safety Gate blocks a run."""


def require_exactly_one_source(input_path):
    """Accept exactly one explicit Excel source file and never scan a directory."""
    if isinstance(input_path, (list, tuple, set)):
        if len(input_path) != 1:
            raise SafetyGateError(
                "SAFETY GATE BLOCKED\nSelect exactly one source workbook."
            )
        input_path = next(iter(input_path))

    path = Path(input_path)

    if not path.exists():
        raise SafetyGateError(
            f"SAFETY GATE BLOCKED\nSource file does not exist:\n{path}"
        )
    if not path.is_file():
        raise SafetyGateError(
            f"SAFETY GATE BLOCKED\nInput must be exactly one file, not a folder:\n{path}"
        )
    if path.suffix.lower() not in {".xlsx", ".xlsm", ".xls"}:
        raise SafetyGateError(
            f"SAFETY GATE BLOCKED\nUnsupported source file:\n{path}"
        )

    return path.resolve()


def load_selected_workbook(input_path, **kwargs):
    """Load only the workbook path accepted by the one-file gate."""
    source_path = require_exactly_one_source(input_path)
    try:
        return load_workbook(source_path, **kwargs)
    except Exception as exc:
        raise SafetyGateError(
            f"SAFETY GATE BLOCKED\nSource workbook could not be loaded:\n{source_path}"
        ) from exc


def read_workbook_structure(input_path):
    """Read worksheet names and header rows from the exact gated workbook."""
    workbook = load_selected_workbook(input_path, data_only=False, read_only=True)
    try:
        structure = []
        for worksheet in workbook.worksheets:
            headers = [
                worksheet.cell(1, column).value
                for column in range(1, worksheet.max_column + 1)
            ]
            structure.append({
                "title": worksheet.title,
                "max_row": worksheet.max_row,
                "max_column": worksheet.max_column,
                "headers": headers,
            })
        return structure
    finally:
        workbook.close()

def norm(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).lower()


def find_header_row(ws, required_terms):
    for r in range(1, min(ws.max_row, 80) + 1):
        vals = [norm(ws.cell(r, c).value) for c in range(1, min(ws.max_column, 15) + 1)]
        joined = " | ".join(vals)
        if all(term in joined for term in required_terms):
            return r
    return None


def parse_amount(value):
    if value is None or value == "":
        return None
    if isinstance(value, (int, float)):
        return float(value)
    s = str(value).replace(",", "").replace("₹", "").strip()
    terms = re.findall(r"-?\d+(?:\.\d+)?", s)
    if not terms:
        return None
    if "+" in s:
        return sum(float(term) for term in terms)
    return float(terms[0])


def normalize_date(value):
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.strftime("%Y-%m-%d")

    text = str(value).strip()
    for date_format in ("%d-%m-%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(text, date_format).strftime("%Y-%m-%d")
        except ValueError:
            continue
    return text


def get_payment_slot(service_date):
    """Return the 10-day payment slot for a milk service date."""
    if isinstance(service_date, datetime):
        service = service_date.date()
    elif isinstance(service_date, date):
        service = service_date
    else:
        normalized = normalize_date(service_date)
        service = datetime.strptime(normalized, "%Y-%m-%d").date()

    if service.day <= 10:
        period = "1-10"
        start = service.replace(day=11)
        end = service.replace(day=20)
    elif service.day <= 20:
        period = "11-20"
        start = service.replace(day=21)
        end = service.replace(day=calendar.monthrange(service.year, service.month)[1])
    else:
        period = "21-EOM"
        if service.month == 12:
            next_month = service.replace(year=service.year + 1, month=1, day=1)
        else:
            next_month = service.replace(month=service.month + 1, day=1)
        start = next_month
        end = next_month.replace(day=10)

    return {
        "service_period": period,
        "payment_slot_start": start.isoformat(),
        "payment_slot_end": end.isoformat(),
    }


def quick_input_key(value):
    if value is None:
        return ""
    return normalize_rule_key(str(value).replace("₹", "").replace("₹", ""))


def quick_input_to_rows(quick_input):
    values = {}
    for raw_key, raw_value in quick_input.items():
        key = quick_input_key(raw_key)
        if key and raw_value not in (None, ""):
            values[key] = raw_value

    rows = []
    milk_rows = []
    date = normalize_date(values.get("date"))
    if not date:
        return rows, milk_rows

    def add_daily_entry(entry, amount, area, category, source_type):
        if amount is None:
            return
        rows.append({
            "DATE": date,
            "AREA": area,
            "CATEGORY": category,
            "ENTRY": entry,
            "AMOUNT (₹)": parse_amount(amount),
            "RAW_AMOUNT": amount,
            "PAYMENT / NOTE": "",
            "SOURCE TYPE": source_type
        })

    add_daily_entry("Sudha Milk", values.get("sudha milk"), "SHOP", "Milk", "SHOP PURCHASE")
    add_daily_entry("Diesel", values.get("diesel"), "KARKHANA / FACTORY", "Karkhana", "KARKHANA EXPENSE")
    add_daily_entry("Bakery", values.get("bakery"), "SHOP", "Shop Purchase", "SHOP PURCHASE")
    add_daily_entry("Water", values.get("water"), "SHOP", "Shop Purchase", "SHOP PURCHASE")
    add_daily_entry("Meal", values.get("meal"), "SHOP", "Shop Purchase", "SHOP PURCHASE")

    sagar_payment = parse_amount(values.get("sagar payment"))
    sagar_litres = parse_amount(values.get("sagar litres"))


    sagar_destination = "KARKHANA / FACTORY"
    sagar_area = "KARKHANA / FACTORY"
    sagar_category = "Karkhana Purchase"
    sagar_source = "KARKHANA PURCHASE"

    if sagar_payment is not None:
        add_daily_entry("Sagar", sagar_payment, sagar_area, sagar_category, sagar_source)
    if sagar_litres is not None:
        milk_rows.append({
            "DATE": date,
            "SUPPLIER": "Sagar",
            "MILK QUANTITY": sagar_litres,
            "UNIT": "Litres",
            "TRIAL / NOTE": "Quick Input",
            "PAYMENT (₹)": sagar_payment
        })

    gokul_payment = parse_amount(values.get("gokul payment"))
    gokul_litres = parse_amount(values.get("gokul litres"))

    gokul_destination = "KARKHANA / FACTORY"
    gokul_area = "KARKHANA / FACTORY"
    gokul_category = "Karkhana Purchase"
    gokul_source = "KARKHANA PURCHASE"


    if gokul_payment is not None:
        add_daily_entry("Gokul", gokul_payment, gokul_area, gokul_category, gokul_source)
    if gokul_litres is not None:
        milk_rows.append({
            "DATE": date,
            "SUPPLIER": "Gokul",
            "MILK QUANTITY": gokul_litres,
            "UNIT": "Litres",
            "TRIAL / NOTE": "Quick Input",
            "PAYMENT (₹)": gokul_payment
        })

    add_daily_entry("Outside Money", values.get("outside money"), "SHOP", "Payment/Control", "OUTSIDE MONEY")
    add_daily_entry("GPay", values.get("gpay"), "SHOP", "Payment/Control", "AXIS BANK GPAY / OUTSIDE MONEY")
    add_daily_entry("LB", values.get("lb"), "SHOP", "Payment/Control", "LB")
    add_daily_entry("SIL", values.get("sil"), "SHOP", "Payment/Control", "SIL")

    return rows, milk_rows


def normalize_rule_key(value):
    if value is None:
        return ""
    return re.sub(r"\s+", " ", str(value).strip()).lower()


def is_duplicate_or_control_row(row):
    entry = normalize_rule_key(row.get("ENTRY"))
    note = normalize_rule_key(row.get("PAYMENT / NOTE"))
    combined = f"{entry} {note}"
    control_markers = (
        "axis bank gpay",
        "payment source",
        "outside money",
        "gpay",
        "duplicate",
        "control row",
        "presentation row",
        "bank gpay"
    )
    return any(marker in combined for marker in control_markers)

def approved_rows_for_summary(rows):
    valid = []
    for row in rows:
        if row.get("DUPLICATE_PRESENTATION_ROW") or row.get("OCR_SCAN_ERROR"):
            continue
        valid.append(row)
    return valid


def approved_expense_rows(rows):
    return [
        row for row in rows
        if not row.get("DUPLICATE_PRESENTATION_ROW")
        and not row.get("OCR_SCAN_ERROR")
        and not row.get("PAYMENT_COLLECTION_ROW")
        and not row.get("MILK_QUANTITY_ROW")
        and not row.get("CONTROL_ROW")
    ]


def _milk_amount_text(row):
    return str(row.get("RAW_AMOUNT", row.get("AMOUNT (₹)")) or "").replace(" ", "")


def detected_milk_supplier(entry, note=""):
    """Detect the supplier only; never use the supplier as the destination."""
    key = normalize_rule_key(entry)
    note_key = normalize_rule_key(note)
    combined = f"{key} {note_key}"
    suppliers = (
        ("Sagar", ("sagar", "sagar milk")),
        ("Gokul", ("gokul", "gokul milk")),
        ("Amul", ("amul", "amul milk")),
        ("Sudha", ("sudha", "sudha milk")),
    )
    for supplier, aliases in suppliers:
        if any(re.search(rf"\b{re.escape(alias)}\b", combined) for alias in aliases):
            return supplier
    return ""


def is_milk_row(row):
    """Recognize milk without assuming Shop/Karkhana from the supplier name."""
    entry_key = normalize_rule_key(row.get("ENTRY"))
    category_key = normalize_rule_key(row.get("CATEGORY"))
    note_key = normalize_rule_key(row.get("PAYMENT / NOTE"))
    source_key = normalize_rule_key(row.get("SOURCE TYPE"))
    return bool(
        detected_milk_supplier(row.get("ENTRY"), row.get("PAYMENT / NOTE"))
        or entry_key == "milk"
        or "milk" in category_key
        or "milk" in note_key
        or "milk" in source_key
    )


def _parse_milk_quantity_expression(value):
    """Parse milk quantity notation: legacy 2/2+2 and full-litre 20/20+20."""
    text = str(value or "").strip().replace(" ", "")
    if not text:
        return None

    # Only treat a simple addition expression as quantity notation.
    if not re.fullmatch(r"\d+(?:\+\d+)?", text):
        return None

    parts = [float(part) for part in text.split("+")]
    if not parts or any(part <= 0 for part in parts):
        return None

    # Guard against interpreting a money amount such as ₹5,483 as litres.
    # Normal milk quantity entries are small daily litre values.
    if any(part > 250 for part in parts) or sum(parts) > 500:
        return None

    # Legacy shorthand: 2 means 20 L and 2+2 means 40 L.
    if all(part == 2 for part in parts):
        return sum(parts) * 10.0

    # From August onward full-litre entries are entered directly:
    # e.g. 20+20 means 40 L.
    return sum(parts)


def is_milk_quantity_row(row):
    if not is_milk_row(row):
        return False
    return _parse_milk_quantity_expression(_milk_amount_text(row)) is not None


def milk_quantity_litres(row):
    return _parse_milk_quantity_expression(_milk_amount_text(row))




def explicit_milk_destination(row):
    area = normalize_rule_key(row.get("AREA"))
    category = normalize_rule_key(row.get("CATEGORY"))
    source_type = normalize_rule_key(row.get("SOURCE TYPE"))
    note = normalize_rule_key(row.get("PAYMENT / NOTE"))

    # Strong destination markers first.
    if "karkhana" in area or "factory" in area:
        return "KARKHANA / FACTORY"
    if "karkhana purchase" in category or "karkhana expense" in category:
        return "KARKHANA / FACTORY"
    if "karkhana purchase" in source_type or "karkhana expense" in source_type:
        return "KARKHANA / FACTORY"
    if re.search(r"\b(karkhana|factory)\b", note):
        return "KARKHANA / FACTORY"

    if area == "shop" or "shop purchase" in category or "shop expense" in category:
        return "SHOP"
    if "shop purchase" in source_type or "shop expense" in source_type:
        return "SHOP"
    if re.search(r"\bshop\b", note):
        return "SHOP"

    return ""


def is_payment_collection_row(row):
    return normalize_rule_key(row.get("ENTRY")) in {
        "gpay",
        "lb",
        "l.b.",
        "l.b. (last balance)",
        "sil",
    }


def authoritative_entry_key(entry):
    """Return a stable alias for diagnostics without any date/value dependency."""
    key = normalize_rule_key(entry)
    aliases = (
        ("sagar", ("sagar", "sagar milk")),
        ("gokul", ("gokul", "gokul milk")),
        ("amul", ("amul", "amul milk")),
        ("sudha milk", ("sudha", "sudha milk")),
        ("bishnu cream", ("bishnu cream", "biscream")),
        ("16 anna", ("16 anna", "16 ana")),
        ("amul axis", ("amul axis",)),
        ("kareem kumar", ("kareem kumar", "varun kumar")),
        ("campa cola", ("campa cola", "campa sale")),
    )
    for canonical, names in aliases:
        if any(name in key for name in names):
            return canonical
    return key or None


def parse_actual_payment_date(row):
    for key in ("PAYMENT DATE", "ACTUAL PAYMENT DATE"):
        if row.get(key):
            return normalize_date(row[key])
    return ""


def payment_slot_diagnostics(rows):
    diagnostics = []
    for row in rows:
        supplier = detected_milk_supplier(row.get("ENTRY"), row.get("PAYMENT / NOTE"))
        if not supplier:
            continue

        service_date = normalize_date(row.get("DATE"))
        slot = get_payment_slot(service_date)
        actual_payment_date = parse_actual_payment_date(row)
        exception = ""
        if actual_payment_date:
            if not slot["payment_slot_start"] <= actual_payment_date <= slot["payment_slot_end"]:
                exception = "PAYMENT_OUTSIDE_EXPECTED_SLOT"

        diagnostics.append({
            "supplier": supplier,
            "service_date": service_date,
            "service_period": slot["service_period"],
            "payment_slot_start": slot["payment_slot_start"],
            "payment_slot_end": slot["payment_slot_end"],
            "actual_payment_date": actual_payment_date or "NOT PROVIDED",
            "amount": parse_amount(row.get("AMOUNT (₹)")),
            "actual_litres": row.get("MILK QUANTITY", ""),
            "exception": exception,
        })
    return diagnostics


def normalize_human_verified_amount(row):
    """Legacy hook retained for compatibility; no date-specific value rewriting."""
    return


def validate_approved_rules(rows, rules):
    """Validate structural safety rules without hard-coding a historical date or totals."""
    errors = []

    for row in rows:
        entry = normalize_rule_key(row.get("ENTRY"))
        amount = parse_amount(row.get("AMOUNT (₹)"))

        # Keep the explicit OCR marker rule only when the source itself identifies
        # the value as an OCR/scan/error/adjustment row. Do not reject a legitimate
        # accounting amount merely because it happens to equal 1,981.
        if amount == 1981 and any(
            marker in entry
            for marker in ("ocr", "scan", "error", "adjustment")
        ):
            row["OCR_SCAN_ERROR"] = True

        if is_milk_row(row):
            destination = explicit_milk_destination(row)
            if not destination:
                errors.append(
                    f"Milk destination is not explicitly recorded for '{row.get('ENTRY', 'Milk')}'. "
                    "Supplier name alone must never decide Shop vs Karkhana."
                )

            if is_milk_quantity_row(row):
                litres = milk_quantity_litres(row)
                if litres is None or litres <= 0:
                    errors.append(
                        f"Milk quantity could not be safely interpreted for '{row.get('ENTRY', 'Milk')}'."
                    )
                else:
                    row["MILK QUANTITY"] = litres
                    row["UNIT"] = "Litres"

    return errors


def extract_quick_input_values(path):
    if not Path(path).exists():
        return {}
    wb = load_selected_workbook(path, data_only=True, read_only=True)
    quick = None
    for ws in wb.worksheets:
        if ws.title.lower() == "quick input":
            quick = ws
            break
    if quick is None:
        wb.close()
        return {}

    mapping = {}
    for row in quick.iter_rows(values_only=True):
        if not row or row[0] in (None, ""):
            continue
        label = str(row[0]).strip()
        value = row[1] if len(row) > 1 else None
        mapping[label] = value
    wb.close()
    return mapping


def extract_canonical_rows(path):
    rows = []

    wb = load_selected_workbook(path, data_only=False, read_only=True)

    for ws in wb.worksheets:
        # Find the canonical header row.
        hr = find_header_row(
            ws,
            ["date", "area", "category", "entry", "amount"]
        )

        if not hr:
            continue

        headers = [
            norm(ws.cell(hr, c).value)
            for c in range(1, ws.max_column + 1)
        ]

        idx = {}

        header_aliases = {
            "date": ["date"],
            "area": ["area"],
            "category": ["category"],
            "entry": ["entry"],
            "amount": ["amount", "amount (₹)"],
            "payment / note": ["payment / note"],
            "payment date": ["payment date", "actual payment date"],
            "payment": ["payment"],
            "source type": ["source type", "source"],
        }

        for key, aliases in header_aliases.items():
            for alias in aliases:
                if alias in headers:
                    idx[key] = headers.index(alias) + 1
                    break

        for values in ws.iter_rows(
            min_row=hr + 1,
            max_row=ws.max_row,
            min_col=1,
            max_col=ws.max_column,
            values_only=True,
        ):
            entry = (
                values[idx["entry"] - 1]
                if "entry" in idx
                else None
            )

            if entry in (None, ""):
                continue

            amount = (
                values[idx["amount"] - 1]
                if "amount" in idx
                else None
            )

            rows.append({
                "DATE": (
                    normalize_date(values[idx["date"] - 1])
                    if "date" in idx
                    else None
                ),
                "AREA": (
                    values[idx["area"] - 1]
                    if "area" in idx
                    else ""
                ),
                "CATEGORY": (
                    values[idx["category"] - 1]
                    if "category" in idx
                    else ""
                ),
                "ENTRY": str(entry).strip(),
                "AMOUNT (₹)": parse_amount(amount),
                "RAW_AMOUNT": amount,
                "PAYMENT / NOTE": (
                    values[idx["payment / note"] - 1]
                    if "payment / note" in idx
                    else None
                ),
                "PAYMENT DATE": (
                    normalize_date(values[idx["payment date"] - 1])
                    if "payment date" in idx
                    else ""
                ),
                "SOURCE TYPE": (
                    values[idx["source type"] - 1]
                    if "source type" in idx
                    else None
                ),
            })

    wb.close()

    return rows

 
def classify_row(row):
    """Apply existing classification rules while preserving explicit milk destination."""
    classified = dict(row)
    entry_key = normalize_rule_key(classified.get("ENTRY"))
    note_key = normalize_rule_key(classified.get("PAYMENT / NOTE"))
    milk_row = is_milk_row(classified)

    normalize_human_verified_amount(classified)

    if entry_key == "diesel":
        classified["AREA"] = "KARKHANA / FACTORY"
        classified["CATEGORY"] = "Karkhana Expense"
        classified["SOURCE TYPE"] = "KARKHANA EXPENSE"

    verified_categories = {
        "bakery": ("SHOP", "Shop Purchase", "SHOP PURCHASE"),
        "water": ("KARKHANA / FACTORY", "Karkhana Expense", "KARKHANA EXPENSE"),
        "16 ana mixture & namkeen": ("SHOP", "Shop Purchase", "SHOP PURCHASE"),
        "varun kumar": ("KARKHANA / FACTORY", "Karkhana Expense", "KARKHANA EXPENSE"),
        "disposable": ("SHOP", "Shop Expense", "SHOP EXPENSE"),
        "chetak": ("SHOP", "Shop Purchase", "SHOP PURCHASE"),
    }
    if entry_key in verified_categories:
        classified["AREA"], classified["CATEGORY"], classified["SOURCE TYPE"] = verified_categories[entry_key]

    if entry_key == "meal":
        classified["AREA"] = "SHOP"
        classified["CATEGORY"] = "Shop Expense"
        classified["SOURCE TYPE"] = "SHOP EXPENSE"
        classified["SPLIT_MEAL"] = True

    # MILK RULE:
    # Supplier (Amul/Sudha/Sagar/Gokul) is metadata, not destination.
    # Explicitly recorded Shop/Karkhana use always wins.
    if milk_row:
        destination = explicit_milk_destination(classified)

        if is_milk_quantity_row(classified):
            classified["MILK QUANTITY"] = milk_quantity_litres(classified)
            classified["UNIT"] = "Litres"
            classified["MILK_QUANTITY_ROW"] = True
            classified["PAYMENT / NOTE"] = (
                classified.get("PAYMENT / NOTE")
                or "Milk quantity notation"
            )

        if destination == "SHOP":
            classified["AREA"] = "SHOP"
            classified["CATEGORY"] = "Milk" if is_milk_quantity_row(classified) else "Shop Purchase"
            classified["SOURCE TYPE"] = "MILK PURCHASE" if is_milk_quantity_row(classified) else "SHOP PURCHASE"
        elif destination == "KARKHANA / FACTORY":
            classified["AREA"] = "KARKHANA / FACTORY"
            classified["CATEGORY"] = "Milk" if is_milk_quantity_row(classified) else "Karkhana Purchase"
            classified["SOURCE TYPE"] = "MILK PURCHASE" if is_milk_quantity_row(classified) else "KARKHANA PURCHASE"
        else:
            # Do not guess. Preserve any source fields and mark for human review.
            classified["MILK_DESTINATION_REVIEW"] = True

    if "outside money" in entry_key or "outside money" in note_key:
        classified["SOURCE TYPE"] = "OUTSIDE MONEY"
        classified["CONTROL_ROW"] = True
        classified["CATEGORY"] = "Payment/Control"

    if "axis bank gpay" in entry_key or "gpay" in entry_key or "bank gpay" in entry_key:
        classified["SOURCE TYPE"] = "AXIS BANK GPAY / OUTSIDE MONEY"
        classified["CONTROL_ROW"] = True
        classified["CATEGORY"] = "Payment/Control"

    if "campa" in entry_key and ("sale" in entry_key or "cola" in entry_key):
        classified["CATEGORY"] = "Shop Purchase"
        classified["AREA"] = "SHOP"
        classified["SOURCE TYPE"] = "SHOP PURCHASE"

    if entry_key == "pitha":
        classified["AREA"] = "SHOP"
        classified["CATEGORY"] = "Shop Purchase"

    if "bishnu cream" in entry_key or "biscream" in entry_key:
        classified["AREA"] = "SHOP"
        classified["CATEGORY"] = "Shop Purchase"
        classified["SOURCE TYPE"] = "SHOP PURCHASE"

    if "suji atta" in entry_key:
        classified["AREA"] = "KARKHANA / FACTORY"
        classified["CATEGORY"] = "Karkhana Purchase"
        classified["SOURCE TYPE"] = "KARKHANA PURCHASE"

    if is_payment_collection_row(classified):
        classified["PAYMENT_COLLECTION_ROW"] = True
        classified["CATEGORY"] = "Payment/Collection"

    if "amul" in entry_key and "axis" in entry_key:
        classified["CATEGORY"] = "Shop Purchase"
        classified["AREA"] = "SHOP"
        classified["SOURCE TYPE"] = "AXIS BANK GPAY / OUTSIDE MONEY"
        classified["PAYMENT / NOTE"] = (
            classified.get("PAYMENT / NOTE")
            or "Axis Bank GPay / Outside Money"
        )

    if is_duplicate_or_control_row(classified) and not ("amul" in entry_key and "axis" in entry_key):
        classified["CONTROL_ROW"] = True
        classified["CATEGORY"] = "Payment/Control"
        if "outside money" not in entry_key and "out money" not in entry_key and "gpay" not in entry_key:
            classified["SOURCE TYPE"] = "CONTROL ROW"

    return classified


def expand_verified_splits(rows):
    # Keep the existing meal split behavior, but do not bind it to 31 July.
    expanded = []
    for row in rows:
        if row.get("SPLIT_MEAL"):
            amount = parse_amount(row.get("AMOUNT (₹)"))
            if amount is not None and amount > 0:
                half = amount / 2
                for area, split_amount in (("SHOP", half), ("KARKHANA / FACTORY", half)):
                    split = dict(row)
                    split["AREA"] = area
                    split["CATEGORY"] = "Shop Expense" if area == "SHOP" else "Karkhana Expense"
                    split["SOURCE TYPE"] = "SHOP EXPENSE" if area == "SHOP" else "KARKHANA EXPENSE"
                    split["ENTRY"] = "Meal (Shop)" if area == "SHOP" else "Meal (Karkhana)"
                    split["AMOUNT (₹)"] = split_amount
                    split["PAYMENT / NOTE"] = "Approved split: Meal divided equally"
                    expanded.append(split)
            else:
                expanded.append(row)
        else:
            expanded.append(row)
    return expanded


def print_approved_diagnostic(rows):
    expense_rows = set(id(row) for row in approved_expense_rows(rows))
    print(
        "DATE | RAW ENTRY | RAW AMOUNT | NORMALIZED ENTRY | SUPPLIER | CATEGORY | "
        "QUANTITY | ACTUAL LITRES | NORMALIZED VALUE | "
        "INCLUDE_IN_EXPENSE_TOTAL | EXCLUSION_REASON"
    )

    for row in rows:
        raw_value = row.get("RAW_AMOUNT", row.get("AMOUNT (₹)"))
        is_quantity = bool(row.get("MILK_QUANTITY_ROW"))
        normalized = parse_amount(row.get("AMOUNT (₹)"))
        normalized_value = "QUANTITY ROW" if is_quantity else (normalized or 0)
        included = id(row) in expense_rows

        if included:
            exclusion_reason = ""
        elif is_quantity:
            exclusion_reason = "MILK QUANTITY ROW"
        elif row.get("PAYMENT_COLLECTION_ROW"):
            exclusion_reason = "PAYMENT/COLLECTION ROW"
        elif row.get("CONTROL_ROW"):
            exclusion_reason = "CONTROL ROW"
        elif row.get("OCR_SCAN_ERROR"):
            exclusion_reason = "OCR/SCAN ERROR"
        elif row.get("DUPLICATE_PRESENTATION_ROW"):
            exclusion_reason = "DUPLICATE/PRESENTATION ROW"
        elif row.get("MILK_DESTINATION_REVIEW"):
            exclusion_reason = "MILK DESTINATION REVIEW"
        else:
            exclusion_reason = "NOT AN APPROVED EXPENSE ROW"

        supplier = detected_milk_supplier(row.get("ENTRY"), row.get("PAYMENT / NOTE"))
        quantity = raw_value if is_quantity else ""
        actual_litres = row.get("MILK QUANTITY", "")

        print(
            f"{normalize_date(row.get('DATE'))} | {row.get('ENTRY')} | {raw_value} | "
            f"{normalize_rule_key(row.get('ENTRY'))} | {supplier} | {row.get('CATEGORY', '')} | "
            f"{quantity} | {actual_litres} | {normalized_value} | "
            f"{included} | {exclusion_reason}"
        )

        if is_quantity and (actual_litres is None or float(actual_litres) <= 0):
            print(
                f"DIAGNOSTIC ALERT: {row.get('ENTRY')} milk quantity is missing or invalid"
            )
        if is_quantity and included:
            print(
                f"DIAGNOSTIC ALERT: {row.get('ENTRY')} quantity row was included in expense total"
            )
        if row.get("MILK_DESTINATION_REVIEW"):
            print(
                f"DIAGNOSTIC ALERT: milk destination is not explicit for {row.get('ENTRY')}; "
                "supplier name is not being used as a destination"
            )


def print_payment_slot_tests():
    test_dates = (
        "2026-08-01", "2026-08-10", "2026-08-11", "2026-08-20",
        "2026-08-21", "2026-08-31", "2026-02-28", "2024-02-29",
        "2026-12-31",
    )
    print("PAYMENT SLOT FUNCTION TESTS")
    for service_date in test_dates:
        slot = get_payment_slot(service_date)
        print(
            f"{service_date} | {slot['service_period']} | "
            f"{slot['payment_slot_start']} through {slot['payment_slot_end']}"
        )


def print_payment_slot_report(rows):
    diagnostics = payment_slot_diagnostics(rows)
    quantity_rows = [row for row in rows if row.get("MILK_QUANTITY_ROW")]
    suppliers = ("Sagar", "Gokul")
    print("PAYMENT SLOT CLASSIFICATION")
    for diagnostic in diagnostics:
        print(
            f"{diagnostic['supplier']} | service={diagnostic['service_date']} | "
            f"period={diagnostic['service_period']} | "
            f"slot={diagnostic['payment_slot_start']}..{diagnostic['payment_slot_end']} | "
            f"actual_payment_date={diagnostic['actual_payment_date']} | "
            f"amount={diagnostic['amount']} | exception={diagnostic['exception'] or 'NONE'}"
        )

    for supplier in suppliers:
        supplier_rows = [
            row for row in quantity_rows
            if detected_milk_supplier(row.get("ENTRY")) == supplier
        ]
        litres = sum(float(row.get("MILK QUANTITY") or 0) for row in supplier_rows)
        period_totals = {"1-10": 0.0, "11-20": 0.0, "21-EOM": 0.0}
        for row in supplier_rows:
            period = get_payment_slot(row.get("DATE"))["service_period"]
            period_totals[period] += float(row.get("MILK QUANTITY") or 0)
        print(f"{supplier} MILK SUPPLY ROWS = {len(supplier_rows)}")
        print(f"{supplier} TOTAL ACTUAL LITRES = {litres:g}")
        print(
            f"{supplier} SERVICE-PERIOD TOTALS = "
            f"1-10:{period_totals['1-10']:g} L, "
            f"11-20:{period_totals['11-20']:g} L, "
            f"21-EOM:{period_totals['21-EOM']:g} L"
        )

    exceptions = [item for item in diagnostics if item["exception"]]
    print(f"PAYMENT_OUTSIDE_EXPECTED_SLOT EXCEPTIONS = {len(exceptions)}")
    for item in exceptions:
        print(f"PAYMENT_OUTSIDE_EXPECTED_SLOT = {item}")


def apply_rules(rows, rules):
    return [classify_row(row) for row in rows]


def make_workbook(rows, output_path,milk_rows=None, *, already_validated=False):
    if not already_validated and validate_approved_rules(rows, {}):
        raise ValueError("Approved accounting rules validation failed before workbook generation.")

    wb = Workbook()
    ws = wb.active
    ws.title = "DAILY ENTRIES"
    ws.append(CANONICAL_HEADERS)

    valid_rows = [
        row for row in approved_rows_for_summary(rows)
        if not row.get("MILK_QUANTITY_ROW")
    ]
    for row in valid_rows:
        ws.append([row.get(h, "") for h in CANONICAL_HEADERS])

    if milk_rows:
        milk = wb.create_sheet("MILK PURCHASE")
        milk.append(MILK_HEADERS)
        for row in milk_rows or []:
            milk.append([row.get(h, "") for h in MILK_HEADERS])
    summary = wb.create_sheet("DAILY SUMMARY")
    summary.append([
        "DATE", "SHOP PURCHASE (₹)", "SHOP EXPENSE (₹)",
        "KARKHANA PURCHASE (₹)", "KARKHANA EXPENSE (₹)", "TOTAL (₹)",
        "OUTSIDE MONEY (₹)", "GPAY (₹)", "LB (₹)", "SIL (₹)",
        "SHOP SALES / CONTROL TOTAL (₹)"
    ])

    dates = sorted({str(r["DATE"]) for r in valid_rows if r.get("DATE") not in (None, "")})
    for i, date in enumerate(dates, start=2):
        summary.cell(i, 1, date)
        for col, source_type in [(2, "SHOP PURCHASE"), (3, "SHOP EXPENSE"), (4, "KARKHANA PURCHASE"), (5, "KARKHANA EXPENSE")]:
            summary.cell(i, col, f'=SUMIFS(\'DAILY ENTRIES\'!$E:$E,\'DAILY ENTRIES\'!$A:$A,$A{i},\'DAILY ENTRIES\'!$G:$G,"{source_type}")')
        summary.cell(i, 6, f"=B{i}+C{i}+D{i}+E{i}")
        summary.cell(i, 7, f'=SUMIFS(\'DAILY ENTRIES\'!$E:$E,\'DAILY ENTRIES\'!$A:$A,$A{i},\'DAILY ENTRIES\'!$G:$G,"OUTSIDE MONEY")+SUMIFS(\'DAILY ENTRIES\'!$E:$E,\'DAILY ENTRIES\'!$A:$A,$A{i},\'DAILY ENTRIES\'!$D:$D,"Amul AXIS")')
        for col, label in [(8, "GPay"), (9, "L.B."), (10, "SIL")]:
            summary.cell(i, col, f'=SUMIFS(\'DAILY ENTRIES\'!$E:$E,\'DAILY ENTRIES\'!$A:$A,$A{i},\'DAILY ENTRIES\'!$D:$D,"{label}")')
        summary.cell(i, 11, f"=H{i}+I{i}+J{i}")

    guide = wb.create_sheet("ENTRY GUIDE")
    guide.append(["FIELD / RULE", "WHAT TO ENTER"])
    guide.append(["DATE", "The slip date"])
    guide.append(["SUDHA MILK", "₹ amount"])
    guide.append(["DIESEL", "₹ amount — automatically classified as KARKHANA / FACTORY"])
    guide.append(["SAGAR / GOKUL", "₹ payment in Daily Entries + litres in Milk Purchase"])
    guide.append(["OUTSIDE MONEY", "Money brought from outside; never sales"])
    guide.append(["GPAY / LB / SIL", "Keep separate from expenses"])
    guide.append(["MILK RULE", "Milk supplier does not decide destination; explicit Shop/Karkhana use wins"])
    guide.append(["MILK QUANTITY", "Legacy 2 = 20 L and 2+2 = 40 L; full entries such as 20+20 = 40 L"])
    guide.append(["GENERIC MILK", "Do not force generic Milk to Shop or Karkhana when destination is not explicitly recorded"])
    guide.append(["AUDIT RULE", "Uncertain values must be reviewed, never silently guessed"])

    quick = wb.create_sheet("QUICK INPUT")
    quick.append(["RADHEY — QUICK DAILY INPUT", ""])
    for label in [
        "DATE", "Sudha Milk ₹", "Diesel ₹", "Bakery ₹", "Water ₹", "Meal ₹",
        "Sagar Milk Payment ₹", "Sagar Milk Litres", "Sagar Milk Destination",
        "Gokul Milk Payment ₹", "Gokul Milk Litres", "Gokul Milk Destination",
        "Outside Money ₹", "GPay ₹", "LB ₹", "SIL ₹"
    ]:
        quick.append([label, ""])

    wb.save(output_path)


def process_folder(input_path, output_dir, archive_dir, exceptions_dir):
    source_path = require_exactly_one_source(input_path)
    output_dir, archive_dir, exceptions_dir = map(Path, [output_dir, archive_dir, exceptions_dir])
    output_dir.mkdir(parents=True, exist_ok=True)
    archive_dir.mkdir(parents=True, exist_ok=True)
    exceptions_dir.mkdir(parents=True, exist_ok=True)

    rules = json.loads(Path("rules.json").read_text(encoding="utf-8"))
    report = []
    try:
        rows = extract_canonical_rows(source_path)
        quick_milk_rows = []
        if not rows:
            quick_values = extract_quick_input_values(source_path)
            if quick_values:
                quick_rows, quick_milk_rows = quick_input_to_rows(quick_values)
                rows = quick_rows
                if quick_milk_rows:
                    rows.extend({
                        "DATE": row["DATE"],
                        "AREA": "SHOP",
                        "CATEGORY": "Milk",
                        "ENTRY": row["SUPPLIER"],
                        "AMOUNT (₹)": row["PAYMENT (₹)"],
                        "PAYMENT / NOTE": "Quick Input",
                        "SOURCE TYPE": "MILK PURCHASE"
                    } for row in quick_milk_rows)

        rows = expand_verified_splits(apply_rules(rows, rules))
        milk_quantity_rows = [
            (row, detected_milk_supplier(row["ENTRY"], row.get("PAYMENT / NOTE")))
            for row in rows
            if row.get("MILK_QUANTITY_ROW")
        ]
        milk_quantity_suppliers = {supplier for _, supplier in milk_quantity_rows}
        first_payment_by_supplier = {}
        for payment in rows:
            if payment.get("MILK_QUANTITY_ROW"):
                continue
            supplier = detected_milk_supplier(payment.get("ENTRY"), payment.get("PAYMENT / NOTE"))
            if supplier in milk_quantity_suppliers and supplier not in first_payment_by_supplier:
                first_payment_by_supplier[supplier] = parse_amount(payment.get("AMOUNT (₹)"))

        quick_milk_rows = [
            {
                "DATE": row["DATE"],
                "SUPPLIER": supplier,
                "MILK QUANTITY": row["MILK QUANTITY"],
                "UNIT": row["UNIT"],
                "TRIAL / NOTE": row.get("PAYMENT / NOTE", ""),
                "PAYMENT (₹)": first_payment_by_supplier.get(supplier, ""),
            }
            for row, supplier in milk_quantity_rows
        ]
        if not rows:
            report.append({"file": source_path.name, "status": "REVIEW", "reason": "No canonical table detected"})
        else:
            print_approved_diagnostic(rows)
            print_payment_slot_report(rows)
            validation_issues = validate_approved_rules(rows, rules)
            if validation_issues:
                wait_for_human_confirmation()
                report.append({"file": source_path.name, "status": "BLOCKED", "reason": "; ".join(validation_issues)})
            else:
                out_name = source_path.stem + "_standardized.xlsx"
                make_workbook(
                    rows,
                    output_dir / out_name,
                    quick_milk_rows,
                    already_validated=True,
                )
                shutil.copy2(source_path, archive_dir / source_path.name)
                report.append({"file": source_path.name, "status": "PROCESSED", "rows": len(rows), "output": out_name})
    except SafetyGateError:
        raise
    except Exception as exc:
        report.append({"file": source_path.name, "status": "ERROR", "reason": str(exc)})

    (exceptions_dir / "processing_report.json").write_text(
        json.dumps(report, indent=2, ensure_ascii=False), encoding="utf-8"
    )
    print(json.dumps(report, indent=2, ensure_ascii=False))


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--input", required=True, help="Exactly one explicit source workbook path")
    ap.add_argument("--output", default="output")
    ap.add_argument("--archive", default="archive")
    ap.add_argument("--exceptions", default="exceptions")
    args = ap.parse_args()
    try:
        print_payment_slot_tests()
        process_folder(args.input, args.output, args.archive, args.exceptions)
    except SafetyGateError as exc:
        print(str(exc))
        raise SystemExit(2)
